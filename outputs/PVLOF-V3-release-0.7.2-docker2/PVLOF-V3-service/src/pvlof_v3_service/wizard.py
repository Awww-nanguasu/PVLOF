"""Interactive Windows-friendly client for the PVLOF-V3 HTTP workflow."""

from __future__ import annotations

import argparse
import csv
import json
import math
import os
import sys
import time
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request, urlopen
from zoneinfo import ZoneInfo


class WizardError(RuntimeError):
    """User-facing HTTP or workflow failure."""


class WizardHTTPError(WizardError):
    """HTTP failure with a machine-readable status code and response detail."""

    def __init__(self, status_code: int, detail: Any):
        self.status_code = int(status_code)
        self.detail = detail
        super().__init__(f"HTTP {self.status_code}: {detail}")


def _http_json(
    base_url: str,
    method: str,
    path: str,
    *,
    api_key: str | None = None,
    payload: dict[str, Any] | None = None,
    timeout: float = 60.0,
) -> dict[str, Any]:
    data = (
        json.dumps(payload, ensure_ascii=False).encode("utf-8")
        if payload is not None
        else None
    )
    headers = {"Accept": "application/json"}
    if payload is not None:
        headers["Content-Type"] = "application/json"
    if api_key:
        headers["X-API-Key"] = api_key
    request = Request(
        base_url.rstrip("/") + "/" + path.lstrip("/"),
        data=data,
        headers=headers,
        method=method.upper(),
    )
    try:
        with urlopen(request, timeout=timeout) as response:  # noqa: S310
            body = response.read()
    except HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        try:
            detail = json.loads(body).get("detail", body)
        except json.JSONDecodeError:
            detail = body
        raise WizardHTTPError(exc.code, detail) from exc
    except (URLError, TimeoutError, OSError) as exc:
        raise WizardError(f"无法访问 PVLOF-V3 服务：{exc}") from exc
    if not body:
        return {}
    try:
        result = json.loads(body)
    except json.JSONDecodeError as exc:
        raise WizardError("服务返回了非 JSON 响应") from exc
    if not isinstance(result, dict):
        raise WizardError("服务响应不是 JSON 对象")
    return result


def _prompt_text(
    label: str,
    *,
    default: str | None = None,
    required: bool = True,
) -> str:
    suffix = f" [默认: {default}]" if default is not None else ""
    while True:
        value = input(f"{label}{suffix}: ").strip()
        if value:
            return value
        if default is not None:
            return default
        if not required:
            return ""
        print("该项不能为空。")


def _prompt_choice(
    label: str,
    choices: list[str],
    *,
    default: int = 1,
) -> int:
    print(label)
    for index, choice in enumerate(choices, start=1):
        print(f"  {index}. {choice}")
    while True:
        raw = input(f"请输入序号 [默认: {default}]: ").strip()
        if not raw:
            return default
        try:
            selected = int(raw)
        except ValueError:
            selected = 0
        if 1 <= selected <= len(choices):
            return selected
        print(f"请输入 1 到 {len(choices)}。")


def _prompt_yes_no(label: str, *, default: bool = True) -> bool:
    hint = "是" if default else "否"
    while True:
        raw = input(f"{label} [默认: {hint}]: ").strip().lower()
        if not raw:
            return default
        if raw in {"y", "yes", "1", "true", "是"}:
            return True
        if raw in {"n", "no", "0", "false", "否"}:
            return False
        print("请输入 y 或 n。")


def _parse_datetime(value: str, timezone_name: str) -> datetime:
    normalized = value.strip().replace("/", "-")
    if len(normalized) == 10:
        normalized += "T00:00:00"
    elif " " in normalized and "T" not in normalized:
        normalized = normalized.replace(" ", "T", 1)
    parsed = datetime.fromisoformat(normalized.replace("Z", "+00:00"))
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=ZoneInfo(timezone_name))
    return parsed


def _format_local_time(value: str | datetime | None, timezone_name: str) -> str:
    """Format a local user-facing time without displaying a UTC offset."""

    if value is None or str(value).strip() == "":
        return "-"
    parsed = (
        value
        if isinstance(value, datetime)
        else _parse_datetime(str(value), timezone_name)
    )
    return parsed.astimezone(ZoneInfo(timezone_name)).strftime(
        "%Y-%m-%d %H:%M:%S"
    )


def _parse_date(value: str) -> date:
    normalized = value.strip().replace("/", "-")
    try:
        return date.fromisoformat(normalized)
    except ValueError as exc:
        raise WizardError("日期格式必须为 YYYY-MM-DD") from exc


def _floor_day(value: datetime) -> datetime:
    return value.replace(hour=0, minute=0, second=0, microsecond=0)


def _ceil_day(value: datetime) -> datetime:
    floored = _floor_day(value)
    return floored if value == floored else floored + timedelta(days=1)


def _recommended_training_window(
    current_range: dict[str, Any],
    timezone_name: str,
) -> tuple[datetime, datetime]:
    minimum = _parse_datetime(
        str(current_range["minimum_time_local"]), timezone_name
    )
    maximum = _parse_datetime(
        str(current_range["maximum_time_local"]), timezone_name
    )
    start = _ceil_day(minimum)
    last_complete_day = _floor_day(maximum)
    end = min(start + timedelta(days=90), last_complete_day)
    if end <= start:
        raise WizardError("当前电站没有完整的可训练日期范围")
    return start, end


def _complete_evaluation_dates(
    current_range: dict[str, Any],
    timezone_name: str,
) -> tuple[date, date]:
    minimum = _parse_datetime(
        str(current_range["minimum_time_local"]), timezone_name
    )
    maximum = _parse_datetime(
        str(current_range["maximum_time_local"]), timezone_name
    )
    minimum_day = _floor_day(minimum)
    first_date = (
        minimum_day.date()
        if minimum == minimum_day
        else (minimum_day + timedelta(days=1)).date()
    )
    last_date = maximum.date()
    if (maximum.hour, maximum.minute, maximum.second) < (23, 55, 0):
        last_date -= timedelta(days=1)
    if last_date < first_date:
        raise WizardError("当前电站没有完整的可评估自然日")
    return first_date, last_date


def _prompt_date_range(
    *,
    default_start: date,
    default_end: date,
    minimum: date,
    maximum: date,
) -> tuple[date, date]:
    print("\n选择评估日期范围（开始日期和结束日期均包含）")
    print(f"可用完整日期：{minimum.isoformat()} 至 {maximum.isoformat()}")
    while True:
        start = _parse_date(
            _prompt_text("开始日期", default=default_start.isoformat())
        )
        end = _parse_date(
            _prompt_text("结束日期", default=default_end.isoformat())
        )
        if end < start:
            print("结束日期不能早于开始日期。")
            continue
        if start < minimum or end > maximum:
            print("选择的日期超出完整电流数据范围。")
            continue
        return start, end


def _prompt_training_date_range(
    *,
    default_start: datetime,
    default_end: datetime,
    minimum: datetime,
    maximum: datetime,
    timezone_name: str,
    minimum_days: int | None = None,
) -> tuple[datetime, datetime]:
    timezone = ZoneInfo(timezone_name)
    first_boundary = _ceil_day(minimum.astimezone(timezone))
    final_boundary = _floor_day(maximum.astimezone(timezone))
    if final_boundary <= first_boundary:
        raise WizardError("当前电站没有完整的可训练自然日")

    first_date = first_boundary.date()
    last_date = (final_boundary - timedelta(days=1)).date()
    default_start_date = default_start.astimezone(timezone).date()
    default_end_date = (
        default_end.astimezone(timezone) - timedelta(days=1)
    ).date()

    print("\n选择训练数据集（开始日期和结束日期均包含）")
    print(
        f"可用完整日期：{first_date.isoformat()} 至 {last_date.isoformat()}"
    )
    while True:
        start_date = _parse_date(
            _prompt_text("开始日期", default=default_start_date.isoformat())
        )
        end_date = _parse_date(
            _prompt_text("结束日期", default=default_end_date.isoformat())
        )
        if end_date < start_date:
            print("结束日期不能早于开始日期。")
            continue
        if start_date < first_date or end_date > last_date:
            print("选择的日期超出完整电流数据范围。")
            continue
        selected_days = (end_date - start_date).days + 1
        if minimum_days is not None and selected_days < minimum_days:
            print(f"训练范围至少需要 {minimum_days} 天。")
            continue

        start = _parse_datetime(start_date.isoformat(), timezone_name)
        end = _parse_datetime(
            (end_date + timedelta(days=1)).isoformat(), timezone_name
        )
        return start, end


def _print_plants(
    plants: list[dict[str, Any]], timezone_name: str
) -> None:
    print("\n可用电站：")
    for index, plant in enumerate(plants, start=1):
        print(
            f"  {index}. 电站 {plant.get('plant_id')} | "
            f"设备 {plant.get('device_count')} | "
            f"文档 {plant.get('documents')} | "
            f"{_format_local_time(plant.get('minimum_time_local'), timezone_name)}"
            " -> "
            f"{_format_local_time(plant.get('maximum_time_local'), timezone_name)}"
        )


def _write_csv(
    path: Path,
    records: list[dict[str, Any]],
    *,
    fallback_columns: list[str],
) -> None:
    columns: list[str] = []
    for record in records:
        for key in record:
            if key not in columns:
                columns.append(key)
    if not columns:
        columns = fallback_columns
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(records)


FORMAL_EVENT_COLUMNS = [
    "事件ID",
    "电站",
    "逆变器",
    "告警组串",
    "正式告警时间点数",
    "正式告警持续时长(分钟)",
    "组串累计异常状态",
    "正式告警起始时间",
    "正式告警终止时间",
]


ALERT_DETAIL_COLUMNS = [
    "row",
    "电站",
    "逆变器",
    "时间节点",
    "当前候选组串",
    "当前告警组串",
    *(f"组串{number:02d}电流(A)" for number in range(1, 31)),
]


def _mapping_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dict):
        return json.dumps(
            value,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        )
    return str(value)


def _formal_event_table(
    events: list[dict[str, Any]],
    timezone_name: str,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for event in events:
        rows.append(
            {
                "事件ID": event.get("event_id", ""),
                "电站": event.get("plant_id", ""),
                "逆变器": event.get("device_no", ""),
                "告警组串": event.get("alert_strings", ""),
                "正式告警时间点数": event.get("time_points", ""),
                "正式告警持续时长(分钟)": event.get(
                    "duration_minutes", ""
                ),
                "组串累计异常状态": _mapping_text(
                    event.get("member_cumulative_anomaly_minutes")
                ),
                "正式告警起始时间": _format_local_time(
                    event.get("raise_time_local") or event.get("raise_time"),
                    timezone_name,
                ),
                "正式告警终止时间": _format_local_time(
                    event.get("end_time_local") or event.get("end_time"),
                    timezone_name,
                ),
            }
        )
    return rows


def _event_id_for_alert_point(
    point: dict[str, Any],
    events: list[dict[str, Any]],
    timezone_name: str,
) -> str:
    """Find the formal member-segment event containing one alert point."""

    point_time_value = point.get("event_time_local") or point.get("event_time")
    if not point_time_value:
        return ""
    point_time = _parse_datetime(str(point_time_value), timezone_name)
    point_key = (
        str(point.get("plant_id", "")),
        str(point.get("device_no", "")),
        str(point.get("alert_strings", "")),
    )
    for event in events:
        event_key = (
            str(event.get("plant_id", "")),
            str(event.get("device_no", "")),
            str(event.get("alert_strings", "")),
        )
        if event_key != point_key:
            continue
        start_value = event.get("raise_time_local") or event.get("raise_time")
        end_value = event.get("end_time_local") or event.get("end_time")
        if not start_value or not end_value:
            continue
        start_time = _parse_datetime(str(start_value), timezone_name)
        end_time = _parse_datetime(str(end_value), timezone_name)
        if start_time <= point_time <= end_time:
            return str(event.get("event_id", ""))
    return ""


def _alert_detail_table(
    points: list[dict[str, Any]],
    events: list[dict[str, Any]],
    timezone_name: str,
) -> list[dict[str, Any]]:
    """Show disjoint current candidates and alerts for each formal event."""

    event_by_id = {
        str(event.get("event_id", "")): (order, event)
        for order, event in enumerate(events, start=1)
    }
    identified_points = [
        (
            point,
            str(point.get("event_id") or _event_id_for_alert_point(
                point, events, timezone_name
            )),
        )
        for point in points
    ]
    def detail_sort_key(item: tuple[dict[str, Any], str]) -> tuple[str, int, int, str]:
        point, event_id = item
        event_order, event = event_by_id.get(event_id, (len(events) + 1, {}))
        try:
            source_row = int(event["row"])
        except (KeyError, TypeError, ValueError):
            source_row = event_order
        local_time = _format_local_time(
            point.get("event_time_local") or point.get("event_time"),
            timezone_name,
        )
        # Source row numbers restart for each local-day batch.
        return (local_time[:10], source_row, event_order, local_time)

    identified_points.sort(key=detail_sort_key)
    rows: list[dict[str, Any]] = []
    for point, event_id in identified_points:
        event_order, event = event_by_id.get(event_id, (None, {}))
        raw_candidates = [
            member.strip()
            for member in str(point.get("candidate_strings") or "").split(",")
            if member.strip()
        ]
        alert_strings = str(point.get("alert_strings") or "")
        alert_members = {
            member.strip() for member in alert_strings.split(",") if member.strip()
        }
        row = {
            # Source row is assigned per local-day batch by the detector.
            "row": event.get("row") or event_order or "",
            "电站": point.get("plant_id", ""),
            "逆变器": point.get("device_no", ""),
            "时间节点": _format_local_time(
                point.get("event_time_local") or point.get("event_time"),
                timezone_name,
            ),
            "当前候选组串": ",".join(
                member for member in raw_candidates if member not in alert_members
            ),
            "当前告警组串": alert_strings,
        }
        for number in range(1, 31):
            value = point.get(f"string_current_{number:02d}")
            try:
                numeric = float(value)
            except (TypeError, ValueError):
                display: float | str = ""
            else:
                display = round(numeric, 2) if math.isfinite(numeric) else ""
            row[f"组串{number:02d}电流(A)"] = display
        rows.append(row)
    return rows


def _write_alert_detail_workbook(
    path: Path,
    records: list[dict[str, Any]],
) -> None:
    """Write the human-review detail table with event identity cells merged."""

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - dependency is installation-tested
        raise WizardError(
            "Excel review export requires openpyxl in the service environment."
        ) from exc

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "告警事件明细"
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "G2"
    worksheet.print_title_rows = "1:1"

    worksheet.append(ALERT_DETAIL_COLUMNS)
    time_column = ALERT_DETAIL_COLUMNS.index("时间节点") + 1
    for record in records:
        values = [record.get(column, "") for column in ALERT_DETAIL_COLUMNS]
        local_time = values[time_column - 1]
        if isinstance(local_time, str) and local_time not in {"", "-"}:
            try:
                values[time_column - 1] = datetime.strptime(
                    local_time,
                    "%Y-%m-%d %H:%M:%S",
                )
            except ValueError:
                pass
        worksheet.append(values)

    font_name = "Microsoft YaHei"
    white_fill = PatternFill("solid", fgColor="FFFFFFFF")
    thin_black = Side(style="thin", color="FF000000")
    cell_border = Border(
        left=thin_black,
        right=thin_black,
        top=thin_black,
        bottom=thin_black,
    )

    for cell in worksheet[1]:
        cell.fill = white_fill
        cell.border = cell_border
        cell.font = Font(
            name=font_name,
            size=10,
            bold=True,
            color="FF000000",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
    worksheet.row_dimensions[1].height = 32

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.fill = white_fill
            cell.border = cell_border
            cell.font = Font(name=font_name, size=10, color="FF000000")
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=False,
            )
        row[time_column - 1].number_format = "yyyy-mm-dd hh:mm:ss"

    current_start_column = ALERT_DETAIL_COLUMNS.index("组串01电流(A)") + 1
    for row in worksheet.iter_rows(
        min_row=2,
        min_col=current_start_column,
        max_col=len(ALERT_DETAIL_COLUMNS),
    ):
        for cell in row:
            cell.number_format = "0.00"

    first_data_row = 2
    last_data_row = len(records) + 1
    row_number = first_data_row
    while row_number <= last_data_row:
        first_time = worksheet.cell(row_number, time_column).value
        identity = (
            *(worksheet.cell(row_number, column).value for column in range(1, 4)),
            first_time.date() if isinstance(first_time, datetime) else str(first_time)[:10],
        )
        group_end = row_number
        while group_end < last_data_row:
            next_time = worksheet.cell(group_end + 1, time_column).value
            following_identity = (
                *(worksheet.cell(group_end + 1, column).value for column in range(1, 4)),
                next_time.date() if isinstance(next_time, datetime) else str(next_time)[:10],
            )
            if following_identity != identity:
                break
            group_end += 1

        if identity[0] not in {None, ""}:
            for column in range(1, 4):
                if group_end > row_number:
                    worksheet.merge_cells(
                        start_row=row_number,
                        start_column=column,
                        end_row=group_end,
                        end_column=column,
                    )
                worksheet.cell(row_number, column).alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )

        row_number = group_end + 1

    widths = {
        1: 8,
        2: 10,
        3: 18,
        4: 21,
        5: 24,
        6: 24,
    }
    for column in range(current_start_column, len(ALERT_DETAIL_COLUMNS) + 1):
        widths[column] = 14
    for column, width in widths.items():
        worksheet.column_dimensions[get_column_letter(column)].width = width

    workbook.save(path)


def _save_evaluation(
    output_directory: Path,
    plant_id: str,
    start_date: date,
    end_date: date,
    timezone_name: str,
    result: dict[str, Any],
) -> dict[str, Path]:
    stamp = f"{start_date:%Y%m%d}_{end_date:%Y%m%d}"
    evaluation_directory = output_directory / f"{plant_id}_{stamp}"
    events_path = evaluation_directory / "正式告警表.csv"
    details_path = evaluation_directory / "告警事件明细.csv"
    details_excel_path = evaluation_directory / "告警事件明细.xlsx"
    metadata_path = evaluation_directory / "metadata.json"
    events = list(result.get("confirmed_events") or [])
    points = list(
        result.get("alert_event_details")
        or result.get("final_alert_points")
        or []
    )
    _write_csv(
        events_path,
        _formal_event_table(events, timezone_name),
        fallback_columns=FORMAL_EVENT_COLUMNS,
    )
    detail_rows = _alert_detail_table(points, events, timezone_name)
    _write_csv(
        details_path,
        detail_rows,
        fallback_columns=ALERT_DETAIL_COLUMNS,
    )
    _write_alert_detail_workbook(details_excel_path, detail_rows)
    metadata = dict(result.get("metadata") or {})
    metadata["export"] = {
        "schema_version": "pvlof-v3-chinese-review-v4",
        "directory": evaluation_directory.name,
        "timezone": timezone_name,
        "date_range_inclusive": {
            "start_date": start_date.isoformat(),
            "end_date": end_date.isoformat(),
        },
        "files": {
            "formal_events": events_path.name,
            "alert_event_details": details_path.name,
            "alert_event_details_excel": details_excel_path.name,
            "metadata": metadata_path.name,
        },
        "alert_event_detail_scope": (
            "two_preceding_confirmation_slots_through_formal_segment_end"
        ),
        "alert_event_detail_current_columns": 30,
        "excel_merged_columns": ALERT_DETAIL_COLUMNS[:3],
        "detail_row_scope": "source_row_number_per_local_day",
        "current_candidate_definition": (
            "raw_candidate_strings_excluding_current_formal_alert_strings"
        ),
    }
    metadata_path.parent.mkdir(parents=True, exist_ok=True)
    metadata_path.write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return {
        "正式告警表": events_path.resolve(),
        "告警事件明细": details_path.resolve(),
        "告警事件明细Excel": details_excel_path.resolve(),
        "元数据": metadata_path.resolve(),
    }


def _preview_events(
    events: list[dict[str, Any]],
    timezone_name: str,
    maximum: int = 20,
) -> None:
    if not events:
        print("本次范围没有正式低电流事件。")
        return
    print(f"\n正式事件预览（前 {min(maximum, len(events))} 条）：")
    for event in events[:maximum]:
        print(
            "  "
            f"{event.get('plant_id')} | {event.get('device_no')} | "
            f"{event.get('alert_strings')} | "
            f"{_format_local_time(event.get('raise_time_local', event.get('raise_time')), timezone_name)}"
            " -> "
            f"{_format_local_time(event.get('end_time_local', event.get('end_time')), timezone_name)}"
        )


def _evaluate_date_range(
    *,
    base_url: str,
    api_key: str | None,
    data_source_id: str,
    calibration_id: str,
    plant_id: str,
    start_date: date,
    end_date: date,
    timeout: float,
) -> dict[str, Any]:
    day_count = (end_date - start_date).days + 1
    all_events: list[dict[str, Any]] = []
    all_points: list[dict[str, Any]] = []
    all_details: list[dict[str, Any]] = []
    daily_metadata: list[dict[str, Any]] = []
    warnings: list[Any] = []
    completed_days = 0
    skipped_dates: list[str] = []
    for index in range(day_count):
        selected_date = start_date + timedelta(days=index)
        print(
            f"[{index + 1}/{day_count}] 正在评估 {selected_date.isoformat()}……"
        )
        try:
            daily = _http_json(
                base_url,
                "POST",
                "/api/v1/evaluations",
                api_key=api_key,
                payload={
                    "data_source_id": data_source_id,
                    "calibration_id": calibration_id,
                    "plant_id": int(plant_id),
                    "start_date": selected_date.isoformat(),
                    "end_date": selected_date.isoformat(),
                },
                timeout=timeout,
            )
        except WizardHTTPError as exc:
            no_current_data = (
                exc.status_code == 404
                and "No calibrated-device current data was found"
                in str(exc.detail)
            )
            if not no_current_data:
                raise
            selected_date_text = selected_date.isoformat()
            warning = f"{selected_date_text}: 无组串电流数据，已跳过"
            skipped_dates.append(selected_date_text)
            warnings.append(warning)
            daily_metadata.append(
                {
                    "request": {
                        "plant_id": int(plant_id),
                        "start_date": selected_date_text,
                        "end_date": selected_date_text,
                    },
                    "execution": {
                        "mode": "daily_independent_batches",
                        "requested_days": 1,
                        "completed_days": 0,
                    },
                    "outputs": {
                        "final_alert_points": 0,
                        "confirmed_events": 0,
                        "alert_event_details": 0,
                    },
                    "warnings": [warning],
                    "daily_batches": [
                        {
                            "date": selected_date_text,
                            "status": "no_current_data",
                            "detail": str(exc.detail),
                        }
                    ],
                }
            )
            print(
                f"[{index + 1}/{day_count}] 跳过：{selected_date_text} "
                "无组串电流数据。"
            )
            continue
        daily_events = list(daily.get("confirmed_events") or [])
        daily_points = list(daily.get("final_alert_points") or [])
        daily_details = list(daily.get("alert_event_details") or [])
        all_events.extend(daily_events)
        all_points.extend(daily_points)
        all_details.extend(daily_details)
        metadata = dict(daily.get("metadata") or {})
        daily_metadata.append(metadata)
        warnings.extend(metadata.get("warnings") or [])
        completed_days += 1
        print(
            f"[{index + 1}/{day_count}] 完成：正式事件 {len(daily_events)} 条，"
            f"正式告警点 {len(daily_points)} 条。"
        )

    if completed_days == 0:
        raise WizardError(
            "所选日期范围内所有日期均无组串电流数据，无法执行评估："
            f"{start_date.isoformat()} 至 {end_date.isoformat()}"
        )

    all_events.sort(
        key=lambda item: (
            str(item.get("raise_time") or item.get("raise_time_local") or ""),
            str(item.get("plant_id") or ""),
            str(item.get("device_no") or ""),
        )
    )
    all_points.sort(
        key=lambda item: (
            str(item.get("event_time") or item.get("event_time_local") or ""),
            str(item.get("plant_id") or ""),
            str(item.get("device_no") or ""),
        )
    )
    all_details.sort(
        key=lambda item: (
            str(item.get("event_time") or item.get("event_time_local") or ""),
            str(item.get("plant_id") or ""),
            str(item.get("device_no") or ""),
            str(item.get("event_id") or ""),
        )
    )
    return {
        "metadata": {
            "request": {
                "plant_id": int(plant_id),
                "start_date": start_date.isoformat(),
                "end_date": end_date.isoformat(),
            },
            "execution": {
                "mode": "daily_independent_batches",
                "requested_days": day_count,
                "completed_days": completed_days,
                "skipped_days": len(skipped_dates),
                "skipped_dates": skipped_dates,
                "history_lookback_hours": 0,
            },
            "outputs": {
                "final_alert_points": len(all_points),
                "confirmed_events": len(all_events),
                "alert_event_details": len(all_details),
            },
            "warnings": warnings,
            "daily_batches": daily_metadata,
        },
        "final_alert_points": all_points,
        "confirmed_events": all_events,
        "alert_event_details": all_details,
    }


def _connect_production_data_source(
    base_url: str,
    api_key: str | None,
) -> dict[str, Any]:
    print("\n正在访问生产环境 Elasticsearch……")
    return _http_json(
        base_url,
        "POST",
        "/api/v1/data-sources",
        api_key=api_key,
        payload={},
    )


def _wait_for_calibration(
    base_url: str,
    api_key: str | None,
    calibration_id: str,
    poll_seconds: float,
) -> dict[str, Any]:
    stage_labels = {
        "queued": "等待执行",
        "discovering_devices": "发现逆变器设备",
        "planning_current_training_sample": "规划分层训练样本",
        "reading_current_training_data": "读取电流训练数据",
        "reading_optional_weather_data": "读取可选气象数据",
        "inferring_string_inventory": "识别有效组串",
        "fitting_initial_peer_model": "训练初始peer模型",
        "fitting_optional_weather_mapping": "训练可选气象映射",
        "fitting_conditioned_base_model": "训练条件化基础模型",
        "fitting_hierarchical_lof_thresholds": "训练分层LOF阈值",
        "saving_and_validating_artifacts": "保存并验证训练产物",
        "completed": "训练完成",
        "failed": "训练失败",
    }
    previous: tuple[Any, Any] | None = None
    while True:
        record = _http_json(
            base_url,
            "GET",
            f"/api/v1/calibrations/{calibration_id}",
            api_key=api_key,
        )
        state = (record.get("stage"), record.get("progress_percent"))
        if state != previous:
            stage = str(record.get("stage", ""))
            detail = ""
            if stage == "reading_current_training_data":
                rows_read = int(record.get("current_rows_read", 0) or 0)
                expected = int(record.get("expected_current_rows", 0) or 0)
                if expected:
                    detail = f" | 已读取 {rows_read:,}/{expected:,} 行"
            print(
                f"训练进度：{record.get('progress_percent', 0)}% | "
                f"{stage_labels.get(stage, stage)}{detail}"
            )
            previous = state
        status = record.get("status")
        if status == "completed":
            return record
        if status == "failed":
            raise WizardError(
                "训练失败："
                f"{record.get('error_type')}: {record.get('error')}"
            )
        time.sleep(poll_seconds)


def _select_or_create_calibration(
    base_url: str,
    api_key: str | None,
    data_source_id: str,
    plant_id: str,
    current_range: dict[str, Any],
    timezone_name: str,
    poll_seconds: float,
) -> tuple[str, datetime]:
    query = urlencode(
        {"data_source_id": data_source_id, "plant_id": plant_id}
    )
    listed = _http_json(
        base_url,
        "GET",
        f"/api/v1/calibrations?{query}",
        api_key=api_key,
    )
    records = list(listed.get("calibrations") or [])
    obsolete = [
        item
        for item in records
        if item.get("status") == "completed"
        and not item.get("compatible_with_service", False)
    ]
    if obsolete:
        print(
            f"发现 {len(obsolete)} 个旧版永久训练结果，其组串识别规则已失效，"
            "本次不会复用。"
        )
    compatible = [
        item for item in records if item.get("compatible_with_service", False)
    ]
    completed = [
        item for item in compatible if item.get("status") == "completed"
    ]
    active = [
        item
        for item in compatible
        if item.get("status") in {"queued", "running"}
    ]
    if completed:
        latest = completed[-1]
        if _prompt_yes_no(
            "发现已完成的永久训练结果 "
            f"{latest.get('calibration_id')}，直接使用吗？",
            default=True,
        ):
            training_end = _parse_datetime(
                str(latest["training_end_utc"]), timezone_name
            ).astimezone(ZoneInfo(timezone_name))
            return str(latest["calibration_id"]), training_end
    if active:
        latest = active[-1]
        if _prompt_yes_no(
            "发现正在进行的训练 "
            f"{latest.get('calibration_id')}，继续等待吗？",
            default=True,
        ):
            completed_record = _wait_for_calibration(
                base_url,
                api_key,
                str(latest["calibration_id"]),
                poll_seconds,
            )
            training_end = _parse_datetime(
                str(completed_record["training_end_utc"]), timezone_name
            ).astimezone(ZoneInfo(timezone_name))
            return str(latest["calibration_id"]), training_end

    recommended_start, recommended_end = _recommended_training_window(
        current_range,
        timezone_name,
    )
    minimum = _parse_datetime(
        str(current_range["minimum_time_local"]), timezone_name
    )
    maximum = _parse_datetime(
        str(current_range["maximum_time_local"]), timezone_name
    )
    print("建议使用连续 30 至 90 天且相对稳定的历史数据；系统最低 20 天。")
    training_start, training_end = _prompt_training_date_range(
        default_start=recommended_start,
        default_end=recommended_end,
        minimum=minimum,
        maximum=maximum,
        timezone_name=timezone_name,
        minimum_days=20,
    )
    if not _prompt_yes_no(
        "训练可能持续较长时间，确认开始吗？", default=True
    ):
        raise WizardError("用户取消训练")
    record = _http_json(
        base_url,
        "POST",
        "/api/v1/calibrations",
        api_key=api_key,
        payload={
            "data_source_id": data_source_id,
            "plant_id": int(plant_id),
            "start_time": training_start.isoformat(),
            "end_time": training_end.isoformat(),
        },
    )
    calibration_id = str(record["calibration_id"])
    print(f"训练任务已创建：{calibration_id}")
    completed_record = _wait_for_calibration(
        base_url,
        api_key,
        calibration_id,
        poll_seconds,
    )
    sampling = dict(completed_record.get("current_sampling") or {})
    elapsed = float(completed_record.get("elapsed_seconds", 0.0) or 0.0)
    print(
        "训练完成，产物已经永久保存。"
        f"用时 {elapsed / 60:.1f} 分钟；"
        f"电流训练行 {int(completed_record.get('current_rows', 0) or 0):,}；"
        f"原始文档 {int(sampling.get('available_documents', 0) or 0):,}。"
    )
    return calibration_id, training_end


def run_wizard(args: argparse.Namespace) -> int:
    base_url = args.base_url.rstrip("/")
    api_key = args.api_key or None
    print("=" * 68)
    print("PVLOF-V3 数据库发现、训练与正式事件评估向导")
    print("=" * 68)
    health = _http_json(base_url, "GET", "/healthz", timeout=15)
    print(
        f"服务已连接：{health.get('status')} | "
        f"版本 {health.get('service_version')}"
    )
    data_source = _connect_production_data_source(base_url, api_key)
    data_source_id = str(data_source["data_source_id"])
    print(
        f"生产环境已连接：{data_source.get('cluster_name')} | "
        f"Elasticsearch {data_source.get('elasticsearch_version')}"
    )

    discovered = _http_json(
        base_url,
        "GET",
        f"/api/v1/data-sources/{data_source_id}/plants",
        api_key=api_key,
    )
    plants = list(discovered.get("plants") or [])
    if not plants:
        raise WizardError("当前设备索引中没有可用电站")
    timezone_name = str(discovered.get("timezone", "Asia/Shanghai"))
    print(f"以下时间均按 {timezone_name} 显示和输入，无需填写 UTC 偏移。")
    _print_plants(plants, timezone_name)
    plant_choice = _prompt_choice(
        "选择需要训练或评估的电站：",
        [f"电站 {plant.get('plant_id')}" for plant in plants],
        default=1,
    )
    plant_id = str(plants[plant_choice - 1]["plant_id"])

    coverage = _http_json(
        base_url,
        "GET",
        f"/api/v1/data-sources/{data_source_id}/plants/{plant_id}/range",
        api_key=api_key,
    )
    current_range = dict(coverage["current"])
    weather_range = dict(coverage["weather"])
    timezone_name = str(coverage.get("timezone", timezone_name))
    print("\n数据范围：")
    print(
        "  电流："
        f"{_format_local_time(current_range.get('minimum_time_local'), timezone_name)}"
        " -> "
        f"{_format_local_time(current_range.get('maximum_time_local'), timezone_name)} | "
        f"设备 {current_range.get('devices')} | "
        f"文档 {current_range.get('documents')}"
    )
    print(
        f"  气象：{'可用' if weather_range.get('available') else '不可用'} | "
        f"{_format_local_time(weather_range.get('minimum_time_local'), timezone_name)}"
        " -> "
        f"{_format_local_time(weather_range.get('maximum_time_local'), timezone_name)}"
    )

    calibration_id, training_end = _select_or_create_calibration(
        base_url,
        api_key,
        data_source_id,
        plant_id,
        current_range,
        timezone_name,
        args.poll_seconds,
    )
    if not _prompt_yes_no("现在开始选择评估日期范围吗？", default=True):
        print(f"已完成。以后选择训练 ID：{calibration_id}")
        return 0

    minimum_date, maximum_date = _complete_evaluation_dates(
        current_range,
        timezone_name,
    )
    training_end_date = training_end.astimezone(
        ZoneInfo(timezone_name)
    ).date()
    default_date = min(max(training_end_date, minimum_date), maximum_date)
    evaluation_start, evaluation_end = _prompt_date_range(
        default_start=default_date,
        default_end=default_date,
        minimum=minimum_date,
        maximum=maximum_date,
    )
    day_count = (evaluation_end - evaluation_start).days + 1
    print(
        f"将按 {timezone_name} 自然日执行 {day_count} 个独立批次，"
        "不读取所选日期之前的数据。"
    )
    result = _evaluate_date_range(
        base_url=base_url,
        api_key=api_key,
        data_source_id=data_source_id,
        calibration_id=calibration_id,
        plant_id=plant_id,
        start_date=evaluation_start,
        end_date=evaluation_end,
        timeout=args.evaluation_timeout,
    )
    events = list(result.get("confirmed_events") or [])
    points = list(result.get("final_alert_points") or [])
    details = list(result.get("alert_event_details") or [])
    print(
        f"评估完成：正式事件 {len(events)} 条，"
        f"正式告警点 {len(points)} 条，复核明细 {len(details)} 条。"
    )
    _preview_events(events, timezone_name)
    paths = _save_evaluation(
        Path(args.output_directory),
        plant_id,
        evaluation_start,
        evaluation_end,
        timezone_name,
        result,
    )
    print("\n输出文件：")
    for name, path in paths.items():
        print(f"  {name}: {path}")
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--base-url",
        default="http://127.0.0.1:8000",
        help="PVLOF-V3 service URL",
    )
    parser.add_argument(
        "--api-key",
        default=os.getenv("PVLOF_CLIENT_API_KEY", ""),
        help="X-API-Key value; preferably set PVLOF_CLIENT_API_KEY",
    )
    parser.add_argument(
        "--output-directory",
        default="outputs/pvlof_v3_cli",
    )
    parser.add_argument("--poll-seconds", type=float, default=5.0)
    parser.add_argument("--evaluation-timeout", type=float, default=1200.0)
    return parser


def main() -> None:
    try:
        code = run_wizard(build_parser().parse_args())
    except KeyboardInterrupt:
        print("\n操作已中断。正在运行的服务端训练任务不会被删除。")
        code = 130
    except (WizardError, ValueError) as exc:
        print(f"\n[FAILED] {exc}", file=sys.stderr)
        code = 1
    raise SystemExit(code)


if __name__ == "__main__":
    main()
