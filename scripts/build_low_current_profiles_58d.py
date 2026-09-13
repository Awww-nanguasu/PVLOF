"""Build low-current events and provisional history profiles from 58-day replay.

The script reads V1.7 Improved V2 full-point Parquets in streaming batches,
keeps raw candidates as the event source, and records final-alert evidence
without using it as a hard filter.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_ROOT = PROJECT_ROOT / "src"
if str(SOURCE_ROOT) not in sys.path:
    sys.path.insert(0, str(SOURCE_ROOT))


import pandas as pd
import pyarrow.parquet as pq

from pv_anomaly.low_current_events import (
    build_device_events,
    build_string_events,
    prepare_candidate_points,
)
from pv_anomaly.low_current_profiles import (
    LowCurrentProfileConfig,
    attach_historical_features,
    build_string_profiles,
    build_valid_day_coverage,
    load_config,
)


KEY_COLUMNS = ["plant_id", "device_no", "event_time", "string_no"]
DAILY_KEYS = ["plant_id", "device_no", "string_no", "local_date"]
OPTIONAL_COLUMN_MAP = {
    "string_current": "string_current",
    "expected_current": "expected_current",
    "residual_ratio": "residual_ratio",
    "conditioned_virtual_irradiance": "conditioned_virtual_irradiance",
    "pvlof_v17_improved_relative_drop": "relative_drop",
    "pvlof_v17_improved_absolute_drop": "absolute_drop",
    "pvlof_v17_improved_internal_relative_drop": "internal_relative_drop",
    "pvlof_v17_improved_internal_absolute_drop": "internal_absolute_drop",
    "pvlof_v17_improved_segmentation_group_relative_drop": "group_relative_drop",
    "pvlof_v17_improved_segmentation_raw_candidate": "segmentation_candidate",
    "collective_raw_alert": "collective_candidate",
    "pvlof_v16_raw_anomaly": "v16_candidate",
    "pvlof_v17_improved_segmentation_original_candidate_accepted": (
        "original_segment_candidate"
    ),
    "pvlof_v17_improved_segmentation_rescue_raw_candidate": (
        "fragmented_reference_rescue_candidate"
    ),
    "pvlof_v17_improved_segmentation_partial_next_segment_raw_candidate": (
        "partial_next_segment_candidate"
    ),
    "pvlof_v17_improved_segmentation_small_candidate_rescue_raw_candidate": (
        "small_candidate_rescue_candidate"
    ),
}


def _as_bool(series: pd.Series) -> pd.Series:
    return series.fillna(False).astype(bool)


def _read_summary_contract(
    plant_directory: Path,
    config: LowCurrentProfileConfig,
) -> dict[str, Any]:
    path = plant_directory / "summary.json"
    if not path.exists():
        return {"summary_path": None, "validated": False}
    payload = json.loads(path.read_text(encoding="utf-8"))
    contract = payload.get("contract", {})
    version = (
        payload.get("versions", {})
        .get("v1_7_improved_v2", {})
        .get("version")
    )
    if version and version != config.source_algorithm_version:
        raise ValueError(
            f"{path} reports source version {version!r}, expected "
            f"{config.source_algorithm_version!r}"
        )
    if contract.get("timezone") and contract["timezone"] != config.timezone:
        raise ValueError(
            f"{path} reports timezone {contract['timezone']!r}, expected "
            f"{config.timezone!r}"
        )
    interval = contract.get("interval_minutes")
    if interval is not None and int(interval) != config.interval_minutes:
        raise ValueError(
            f"{path} reports interval {interval}, expected "
            f"{config.interval_minutes}"
        )
    return {
        "summary_path": str(path),
        "validated": True,
        "contract": contract,
        "source_algorithm_version": version,
    }


def _read_plant_points(
    plant_directory: Path,
    config: LowCurrentProfileConfig,
) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, Any]]:
    path = plant_directory / config.source_points_filename
    if not path.exists():
        raise FileNotFoundError(path)
    parquet = pq.ParquetFile(path)
    schema_names = set(parquet.schema_arrow.names)
    critical_map = {
        config.raw_candidate_column: "raw_candidate",
        config.final_alert_column: "final_alert",
        config.valid_point_column: "valid_point",
    }
    required = {*KEY_COLUMNS, *critical_map}
    missing = sorted(required - schema_names)
    if missing:
        raise ValueError(f"{path} is missing required columns: {missing}")

    selected_map = {
        **critical_map,
        **{
            source: destination
            for source, destination in OPTIONAL_COLUMN_MAP.items()
            if source in schema_names
        },
    }
    selected_columns = list(dict.fromkeys([*KEY_COLUMNS, *selected_map]))
    missing_optional = sorted(set(OPTIONAL_COLUMN_MAP) - schema_names)
    candidates: list[pd.DataFrame] = []
    daily_counts: list[pd.DataFrame] = []
    rows_scanned = 0
    valid_points = 0
    raw_candidates = 0
    final_alert_points = 0
    invalid_candidate_points = 0
    null_key_rows = 0
    first_time: pd.Timestamp | None = None
    last_time: pd.Timestamp | None = None

    for batch in parquet.iter_batches(
        batch_size=config.batch_size,
        columns=selected_columns,
    ):
        chunk = batch.to_pandas()
        rows_scanned += len(chunk)
        chunk["plant_id"] = chunk["plant_id"].astype("string")
        chunk["device_no"] = chunk["device_no"].astype("string")
        chunk["string_no"] = pd.to_numeric(
            chunk["string_no"], errors="coerce"
        ).astype("Int64")
        chunk["event_time"] = pd.to_datetime(
            chunk["event_time"], errors="coerce", utc=True
        )
        null_key = chunk[KEY_COLUMNS].isna().any(axis=1)
        null_key_rows += int(null_key.sum())
        valid_key_chunk = chunk[~null_key].copy()

        raw_mask = _as_bool(valid_key_chunk[config.raw_candidate_column])
        final_mask = _as_bool(valid_key_chunk[config.final_alert_column])
        valid_mask = _as_bool(valid_key_chunk[config.valid_point_column])
        valid_points += int(valid_mask.sum())
        raw_candidates += int(raw_mask.sum())
        final_alert_points += int(final_mask.sum())
        invalid_candidate_points += int((raw_mask & ~valid_mask).sum())

        if not valid_key_chunk.empty:
            chunk_first = valid_key_chunk["event_time"].min()
            chunk_last = valid_key_chunk["event_time"].max()
            first_time = (
                chunk_first
                if first_time is None or chunk_first < first_time
                else first_time
            )
            last_time = (
                chunk_last
                if last_time is None or chunk_last > last_time
                else last_time
            )

        valid_rows = valid_key_chunk[valid_mask][
            ["plant_id", "device_no", "string_no", "event_time"]
        ].copy()
        if not valid_rows.empty:
            valid_rows["local_date"] = (
                valid_rows["event_time"]
                .dt.tz_convert(config.timezone)
                .dt.strftime("%Y-%m-%d")
            )
            daily_counts.append(
                valid_rows.groupby(DAILY_KEYS, observed=True)
                .size()
                .rename("valid_point_count")
                .reset_index()
            )

        selected_candidates = valid_key_chunk[raw_mask][selected_columns].copy()
        if not selected_candidates.empty:
            selected_candidates = selected_candidates.rename(columns=selected_map)
            candidates.append(selected_candidates)

    candidate_frame = (
        pd.concat(candidates, ignore_index=True)
        if candidates
        else pd.DataFrame(columns=[*KEY_COLUMNS, *selected_map.values()])
    )
    daily_frame = (
        pd.concat(daily_counts, ignore_index=True)
        if daily_counts
        else pd.DataFrame(columns=[*DAILY_KEYS, "valid_point_count"])
    )
    duplicate_candidate_rows = int(
        candidate_frame.duplicated(KEY_COLUMNS, keep=False).sum()
    )
    if duplicate_candidate_rows:
        raise ValueError(
            f"{path} contains {duplicate_candidate_rows} duplicate raw-candidate keys"
        )

    contract = _read_summary_contract(plant_directory, config)
    report = {
        "plant_directory": str(plant_directory),
        "input": str(path),
        "parquet_rows": int(parquet.metadata.num_rows),
        "rows_scanned": int(rows_scanned),
        "schema_columns": int(len(schema_names)),
        "selected_columns": selected_columns,
        "missing_optional_columns": missing_optional,
        "null_key_rows": null_key_rows,
        "valid_string_points": valid_points,
        "raw_candidate_string_points": raw_candidates,
        "final_alert_string_points": final_alert_points,
        "raw_candidates_not_v17_valid": invalid_candidate_points,
        "duplicate_candidate_rows": duplicate_candidate_rows,
        "first_time_utc": str(first_time) if first_time is not None else None,
        "last_time_utc": str(last_time) if last_time is not None else None,
        "contract": contract,
    }
    return candidate_frame, daily_frame, report


def _prepare_output(path: Path, *, overwrite: bool) -> None:
    if path.exists() and any(path.iterdir()) and not overwrite:
        raise FileExistsError(
            f"Output directory is not empty: {path}. Pass --overwrite to "
            "replace the known profile outputs."
        )
    path.mkdir(parents=True, exist_ok=True)


def _json_ready(value: Any) -> Any:
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, pd.Timestamp):
        return str(value)
    if pd.isna(value):
        return None
    return value


def _excel_value(value: Any) -> Any:
    if value is pd.NA or value is pd.NaT:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, pd.Timestamp):
        return str(value)
    if hasattr(value, "item"):
        try:
            return value.item()
        except ValueError:
            pass
    return value


def _write_sheet(workbook, title: str, frame: pd.DataFrame) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    worksheet = workbook.create_sheet(title=title)
    columns = [str(column) for column in frame.columns]
    worksheet.append(columns)
    for row in frame.itertuples(index=False, name=None):
        worksheet.append([_excel_value(value) for value in row])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.freeze_panes = "A2"
    if columns:
        worksheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(columns))}{max(worksheet.max_row, 1)}"
        )

    sample_limit = min(worksheet.max_row, 500)
    for column_number, column_name in enumerate(columns, 1):
        width = len(column_name)
        for row_number in range(2, sample_limit + 1):
            value = worksheet.cell(row=row_number, column=column_number).value
            width = max(width, len(str(value)) if value is not None else 0)
        worksheet.column_dimensions[get_column_letter(column_number)].width = min(
            max(width + 2, 10),
            38,
        )


def _workbook_frames(
    event_profiles: pd.DataFrame,
    string_profiles: pd.DataFrame,
    device_events: pd.DataFrame,
    audits: list[dict[str, Any]],
    config: LowCurrentProfileConfig,
) -> dict[str, pd.DataFrame]:
    event_drop = {
        "candidate_start_time",
        "candidate_end_time",
        "candidate_window_end_time",
    }
    event_review = event_profiles[
        [column for column in event_profiles.columns if column not in event_drop]
    ].copy()
    preferred = [
        "manual_profile_class",
        "manual_fixed_time",
        "manual_review_note",
        "profile_class",
        "fixed_time_tag",
        "plant_id",
        "device_no",
        "string_no",
        "candidate_start_time_local",
        "candidate_end_time_local",
    ]
    event_review = event_review[
        [column for column in preferred if column in event_review]
        + [column for column in event_review if column not in preferred]
    ]

    profile_review = string_profiles.copy()
    profile_preferred = [
        "manual_profile_class",
        "manual_fixed_time",
        "manual_review_note",
        "profile_class",
        "fixed_time_tag",
        "plant_id",
        "device_no",
        "string_no",
    ]
    profile_review = profile_review[
        [column for column in profile_preferred if column in profile_review]
        + [column for column in profile_review if column not in profile_preferred]
    ]

    class_summary = (
        string_profiles.groupby(
            ["plant_id", "profile_class", "fixed_time_tag"],
            observed=True,
            dropna=False,
        )
        .size()
        .rename("string_profiles")
        .reset_index()
        if not string_profiles.empty
        else pd.DataFrame()
    )
    quality_rows = []
    for audit in audits:
        quality_rows.append(
            {
                key: value
                for key, value in audit.items()
                if key not in {"selected_columns", "missing_optional_columns", "contract"}
            }
            | {
                "selected_columns": ",".join(audit["selected_columns"]),
                "missing_optional_columns": ",".join(
                    audit["missing_optional_columns"]
                ),
            }
        )
    notes = pd.DataFrame(
        {
            "项目": [
                "事件来源",
                "事件合并",
                "事件级分类",
                "回顾性画像",
                "分类边界",
                "归因边界",
            ],
            "说明": [
                "PVLOF V1.7 Improved V2 原始点级候选；最终告警仅作为属性保留。",
                "同一物理组串严格连续的五分钟候选点合并为一个事件。",
                "只使用当前事件开始之前的历史，避免未来数据泄漏。",
                "组串画像使用整个输入周期，仅用于人工复核和汇报。",
                "当前长期/突发阈值均为暂定参数，需人工标注后重新标定。",
                "画像描述历史模式，不直接认定遮挡、积灰或硬件故障原因。",
            ],
        }
    )
    parameters = pd.DataFrame(
        [
            {"参数": key, "值": json.dumps(value, ensure_ascii=False)}
            for key, value in config.to_dict().items()
        ]
    )
    return {
        "说明": notes,
        "事件明细": event_review,
        "组串画像": profile_review,
        "设备事件": device_events.drop(
            columns=[
                "candidate_start_time",
                "candidate_end_time",
                "candidate_window_end_time",
            ],
            errors="ignore",
        ),
        "分类汇总": class_summary,
        "数据质量": pd.DataFrame(quality_rows),
        "参数说明": parameters,
    }


def _write_workbook(
    path: Path,
    event_profiles: pd.DataFrame,
    string_profiles: pd.DataFrame,
    device_events: pd.DataFrame,
    audits: list[dict[str, Any]],
    config: LowCurrentProfileConfig,
) -> None:
    try:
        from openpyxl import Workbook
    except ImportError as error:
        raise RuntimeError(
            "Excel review output requires openpyxl. Other Parquet/CSV outputs "
            "can be generated with --skip-workbook."
        ) from error

    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, frame in _workbook_frames(
        event_profiles,
        string_profiles,
        device_events,
        audits,
        config,
    ).items():
        _write_sheet(workbook, title, frame)
    workbook.save(path)


def _classification_counts(frame: pd.DataFrame) -> dict[str, int]:
    if frame.empty or "profile_class_code" not in frame:
        return {}
    return {
        str(key): int(value)
        for key, value in frame["profile_class_code"]
        .value_counts(dropna=False)
        .sort_index()
        .items()
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--input-root",
        default="artifacts/models/pvlof_58d_replay",
    )
    parser.add_argument("--plant-id", action="append", default=[])
    parser.add_argument(
        "--config",
        default="configs/low_current_profile_v01.json",
    )
    parser.add_argument(
        "--output-directory",
        default="outputs/low_current_profile_v01_20260601_20260728",
    )
    parser.add_argument("--skip-workbook", action="store_true")
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()

    config = load_config(args.config)
    input_root = Path(args.input_root)
    plant_directories = (
        [input_root / f"plant_id={plant_id}" for plant_id in args.plant_id]
        if args.plant_id
        else sorted(input_root.glob("plant_id=*"))
    )
    if not plant_directories:
        raise FileNotFoundError(f"No plant_id=* directories found under {input_root}")
    missing_directories = [
        str(path) for path in plant_directories if not path.is_dir()
    ]
    if missing_directories:
        raise FileNotFoundError(
            f"Requested plant directories do not exist: {missing_directories}"
        )

    output = Path(args.output_directory)
    _prepare_output(output, overwrite=args.overwrite)
    candidate_frames: list[pd.DataFrame] = []
    daily_frames: list[pd.DataFrame] = []
    audits: list[dict[str, Any]] = []
    for plant_directory in plant_directories:
        print(f"Reading {plant_directory}", file=sys.stderr)
        candidates, daily, audit = _read_plant_points(plant_directory, config)
        candidate_frames.append(candidates)
        daily_frames.append(daily)
        audits.append(audit)

    candidate_source = pd.concat(candidate_frames, ignore_index=True)
    daily_source = pd.concat(daily_frames, ignore_index=True)
    candidate_points = prepare_candidate_points(
        candidate_source,
        timezone=config.timezone,
    )
    valid_day_coverage = build_valid_day_coverage(
        daily_source,
        minimum_valid_points_per_day=config.minimum_valid_points_per_day,
    )
    string_events = build_string_events(
        candidate_points,
        timezone=config.timezone,
        interval_minutes=config.interval_minutes,
        entry_consecutive=config.entry_consecutive,
    )
    device_events = build_device_events(
        candidate_points,
        timezone=config.timezone,
        interval_minutes=config.interval_minutes,
    )
    event_profiles = attach_historical_features(
        string_events,
        valid_day_coverage,
        config,
    )
    string_profiles = build_string_profiles(
        event_profiles,
        valid_day_coverage,
        config,
    )

    paths = {
        "candidate_points": output / "low_current_candidate_points.parquet",
        "valid_day_coverage": output / "low_current_valid_day_coverage.parquet",
        "string_events": output / "low_current_string_events.parquet",
        "device_events": output / "low_current_device_events.parquet",
        "event_profiles": output / "low_current_event_profiles.parquet",
        "string_profiles": output / "low_current_string_profiles.parquet",
        "event_review_csv": output / "low_current_event_review.csv",
        "string_profile_csv": output / "low_current_string_profiles.csv",
        "device_event_csv": output / "low_current_device_events.csv",
        "data_quality": output / "data_quality_summary.json",
        "config_snapshot": output / "config_snapshot.json",
        "summary": output / "summary.json",
        "workbook": output / "low_current_profile_review.xlsx",
    }
    candidate_points.to_parquet(paths["candidate_points"], index=False)
    valid_day_coverage.to_parquet(paths["valid_day_coverage"], index=False)
    string_events.to_parquet(paths["string_events"], index=False)
    device_events.to_parquet(paths["device_events"], index=False)
    event_profiles.to_parquet(paths["event_profiles"], index=False)
    string_profiles.to_parquet(paths["string_profiles"], index=False)
    event_profiles.to_csv(paths["event_review_csv"], index=False, encoding="utf-8-sig")
    string_profiles.to_csv(
        paths["string_profile_csv"], index=False, encoding="utf-8-sig"
    )
    device_events.to_csv(
        paths["device_event_csv"], index=False, encoding="utf-8-sig"
    )

    quality_summary = {
        "inputs": audits,
        "aggregate": {
            "rows_scanned": int(sum(item["rows_scanned"] for item in audits)),
            "valid_string_points": int(
                sum(item["valid_string_points"] for item in audits)
            ),
            "raw_candidate_string_points": int(
                sum(item["raw_candidate_string_points"] for item in audits)
            ),
            "final_alert_string_points": int(
                sum(item["final_alert_string_points"] for item in audits)
            ),
            "raw_candidates_not_v17_valid": int(
                sum(item["raw_candidates_not_v17_valid"] for item in audits)
            ),
            "null_key_rows": int(sum(item["null_key_rows"] for item in audits)),
            "duplicate_candidate_rows": int(
                sum(item["duplicate_candidate_rows"] for item in audits)
            ),
        },
    }
    paths["data_quality"].write_text(
        json.dumps(quality_summary, ensure_ascii=False, indent=2, default=_json_ready)
        + "\n",
        encoding="utf-8",
    )
    paths["config_snapshot"].write_text(
        json.dumps(config.to_dict(), ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    if not args.skip_workbook:
        _write_workbook(
            paths["workbook"],
            event_profiles,
            string_profiles,
            device_events,
            audits,
            config,
        )

    summary = {
        "contract": {
            "profile_version": config.version,
            "source_algorithm_version": config.source_algorithm_version,
            "timezone": config.timezone,
            "interval_minutes": config.interval_minutes,
            "plant_ids": sorted(candidate_points["plant_id"].astype(str).unique()),
            "provisional_thresholds": config.provisional_thresholds,
            "event_source": "raw_candidate",
            "final_alert_is_filter": False,
            "causal_event_history": True,
        },
        "counts": {
            "candidate_string_points": int(len(candidate_points)),
            "candidate_device_times": int(
                candidate_points[
                    ["plant_id", "device_no", "event_time"]
                ].drop_duplicates().shape[0]
            ),
            "valid_string_days": int(valid_day_coverage["is_valid_day"].sum()),
            "string_events": int(len(string_events)),
            "device_events": int(len(device_events)),
            "three_point_string_events": int(
                string_events["meets_three_point_rule"].sum()
            ),
            "string_events_with_final_alert": int(
                string_events["has_final_alert"].sum()
            ),
            "event_profile_classes": _classification_counts(event_profiles),
            "retrospective_string_classes": _classification_counts(
                string_profiles
            ),
            "fixed_time_string_profiles": int(
                string_profiles["fixed_time_tag"].sum()
            ),
        },
        "configuration": config.to_dict(),
        "inputs": [str(path) for path in plant_directories],
        "outputs": {
            key: str(value)
            for key, value in paths.items()
            if key != "workbook" or not args.skip_workbook
        },
    }
    paths["summary"].write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, default=_json_ready) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2, default=_json_ready))


if __name__ == "__main__":
    main()
