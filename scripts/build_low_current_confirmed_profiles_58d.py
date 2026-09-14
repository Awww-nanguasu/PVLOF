"""Build low-current profiles from confirmed V1.7 Improved V2 events."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_ROOT = PROJECT_ROOT / "src"
if str(SOURCE_ROOT) not in sys.path:
    sys.path.insert(0, str(SOURCE_ROOT))


import pandas as pd

from pv_anomaly.low_current_confirmed_profiles import (
    ConfirmedLowCurrentProfileConfig,
    attach_candidate_metrics,
    build_workbook_crosswalk,
    load_confirmed_config,
    load_confirmed_workbook,
    load_replay_confirmed_tables,
    prepare_unconfirmed_runs,
    read_parquet_file,
)
from pv_anomaly.low_current_profiles import (
    attach_historical_features,
    build_string_profiles,
)


def _prepare_output(path: Path, *, overwrite: bool) -> None:
    if path.exists() and any(path.iterdir()) and not overwrite:
        raise FileExistsError(
            f"Output directory is not empty: {path}. Pass --overwrite to "
            "replace the known v0.2 outputs."
        )
    path.mkdir(parents=True, exist_ok=True)


def _json_ready(value: Any) -> Any:
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, pd.Timestamp):
        return str(value)
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if hasattr(value, "item"):
        try:
            return value.item()
        except ValueError:
            pass
    return value


def _excel_value(value: Any) -> Any:
    if value is pd.NA or value is pd.NaT:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, (list, tuple, set)):
        return ",".join(str(item) for item in value)
    if isinstance(value, pd.Timestamp):
        return str(value)
    if hasattr(value, "item"):
        try:
            return value.item()
        except ValueError:
            pass
    return value


def _write_sheet(workbook, title: str, frame: pd.DataFrame) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    worksheet = workbook.create_sheet(title=title)
    columns = [str(column) for column in frame.columns]
    worksheet.append(columns)
    for row in frame.itertuples(index=False, name=None):
        worksheet.append([_excel_value(value) for value in row])

    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.freeze_panes = "A2"
    if columns:
        worksheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(columns))}{max(worksheet.max_row, 1)}"
        )
    for column_number, column_name in enumerate(columns, 1):
        width = len(column_name)
        for row_number in range(2, min(worksheet.max_row, 500) + 1):
            value = worksheet.cell(row=row_number, column=column_number).value
            width = max(width, len(str(value)) if value is not None else 0)
        worksheet.column_dimensions[get_column_letter(column_number)].width = min(
            max(width + 2, 10), 42
        )


def _review_columns(frame: pd.DataFrame, preferred: list[str]) -> pd.DataFrame:
    utc_columns = {
        column
        for column in frame.columns
        if column.endswith("_time") and f"{column}_local" in frame.columns
    }
    result = frame.drop(columns=sorted(utc_columns), errors="ignore").copy()
    order = [column for column in preferred if column in result]
    return result[order + [column for column in result if column not in order]]


def _classification_summary(
    event_profiles: pd.DataFrame,
    string_profiles: pd.DataFrame,
) -> pd.DataFrame:
    rows: list[pd.DataFrame] = []
    for level, frame, count_name in (
        ("事件级", event_profiles, "记录数"),
        ("组串级", string_profiles, "记录数"),
    ):
        if frame.empty:
            continue
        grouped = (
            frame.groupby(
                ["plant_id", "profile_class", "fixed_time_tag"],
                observed=True,
                dropna=False,
            )
            .size()
            .rename(count_name)
            .reset_index()
        )
        grouped.insert(0, "统计层级", level)
        rows.append(grouped)
    return pd.concat(rows, ignore_index=True) if rows else pd.DataFrame()


def _workbook_frames(
    *,
    event_profiles: pd.DataFrame,
    string_profiles: pd.DataFrame,
    device_events: pd.DataFrame,
    unconfirmed: pd.DataFrame,
    crosswalk: pd.DataFrame,
    replay_audits: list[dict[str, Any]],
    crosswalk_audit: dict[str, Any],
    metric_audit: dict[str, Any],
    config: ConfirmedLowCurrentProfileConfig,
) -> dict[str, pd.DataFrame]:
    event_review = _review_columns(
        event_profiles,
        [
            "manual_profile_class",
            "manual_fixed_time",
            "manual_review_note",
            "profile_class",
            "fixed_time_tag",
            "plant_id",
            "device_no",
            "string_no",
            "candidate_start_time_local",
            "alert_confirm_time_local",
            "candidate_end_time_local",
            "event_end_time_local",
            "string_event_id",
            "source_event_key",
            "comparison_event_ids",
        ],
    )
    string_review = _review_columns(
        string_profiles,
        [
            "manual_profile_class",
            "manual_fixed_time",
            "manual_review_note",
            "profile_class",
            "fixed_time_tag",
            "plant_id",
            "device_no",
            "string_no",
        ],
    )
    notes = pd.DataFrame(
        {
            "项目": [
                "正式事件入口",
                "精确事件边界",
                "候选点用途",
                "历史统计口径",
                "时间字段",
                "分类边界",
                "归因边界",
            ],
            "说明": [
                "仅使用 pvlof_58d_confirmed_events.xlsx 中 PVLOF_V1_7_IMPROVED_V2 非空的事件。",
                "事件计数与起止边界来自 V2 events/evidence Parquet；跨多个对比表片段仍只计一次。",
                "候选点只补充跌幅、形态和电流等证据，不产生画像事件，也不进入历史频次。",
                "事件级标签只看当前事件开始前的正式事件；组串画像为整个 58 天周期的回顾统计。",
                "candidate_start 为异常起点，alert_confirm 为正式确认点，event_end 为恢复或结束点。",
                "长期重复型/突发型阈值仍为暂定参数，需结合人工标签继续标定。",
                "本模型描述告警历史画像，不直接断言遮挡、积灰、老化或硬件故障原因。",
            ],
        }
    )
    quality = pd.DataFrame(
        [
            *[{"检查范围": "源文件", **item} for item in replay_audits],
            {"检查范围": "正式事件映射", **crosswalk_audit},
            {"检查范围": "候选指标补充", **metric_audit},
        ]
    )
    parameters = pd.DataFrame(
        [
            {"参数": key, "值": json.dumps(value, ensure_ascii=False)}
            for key, value in config.to_dict().items()
        ]
    )
    return {
        "说明": notes,
        "正式事件画像": event_review,
        "组串画像": string_review,
        "正式设备事件": _review_columns(device_events, []),
        "未确认候选": _review_columns(unconfirmed, []),
        "映射明细": _review_columns(crosswalk, []),
        "分类汇总": _classification_summary(event_profiles, string_profiles),
        "数据质量": quality,
        "参数": parameters,
    }


def _write_workbook(path: Path, frames: dict[str, pd.DataFrame]) -> None:
    try:
        from openpyxl import Workbook
    except ImportError as error:
        raise RuntimeError(
            "Excel output requires openpyxl>=3.1. Use --skip-workbook only "
            "when Parquet/CSV outputs are sufficient."
        ) from error
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, frame in frames.items():
        _write_sheet(workbook, title, frame)
    workbook.save(path)


def _classification_counts(frame: pd.DataFrame) -> dict[str, int]:
    if frame.empty or "profile_class_code" not in frame:
        return {}
    return {
        str(key): int(value)
        for key, value in frame["profile_class_code"]
        .value_counts(dropna=False)
        .sort_index()
        .items()
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--confirmed-events-workbook",
        default="outputs/pvlof_58d_confirmed_events.xlsx",
    )
    parser.add_argument(
        "--input-root",
        default="artifacts/models/pvlof_58d_replay",
    )
    parser.add_argument(
        "--candidate-profile-directory",
        default="outputs/low_current_profile_v01_20260601_20260728",
    )
    parser.add_argument(
        "--config",
        default="configs/low_current_profile_v02.json",
    )
    parser.add_argument(
        "--output-directory",
        default="outputs/low_current_profile_v02_confirmed_20260601_20260728",
    )
    parser.add_argument("--plant-id", action="append", default=[])
    parser.add_argument("--skip-workbook", action="store_true")
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()

    config = load_confirmed_config(args.config)
    output = Path(args.output_directory)
    _prepare_output(output, overwrite=args.overwrite)

    device_events, memberships, workbook_rows = load_confirmed_workbook(
        args.confirmed_events_workbook,
        config,
    )
    if args.plant_id:
        requested_plants = {str(value) for value in args.plant_id}
        device_events = device_events[
            device_events["plant_id"].astype(str).isin(requested_plants)
        ].reset_index(drop=True)
        memberships = memberships[
            memberships["plant_id"].astype(str).isin(requested_plants)
        ].reset_index(drop=True)
        if memberships.empty:
            raise ValueError(
                "The confirmed-event workbook has no V2 memberships for "
                f"the requested plants: {sorted(requested_plants)}"
            )
    source_events, evidence, unconfirmed_source, replay_audits = (
        load_replay_confirmed_tables(
            args.input_root,
            config,
            plant_ids=args.plant_id,
        )
    )
    confirmed_events, crosswalk, crosswalk_audit = build_workbook_crosswalk(
        source_events,
        evidence,
        memberships,
        strict=config.strict_workbook_match,
    )

    candidate_directory = Path(args.candidate_profile_directory)
    candidate_path = candidate_directory / config.candidate_points_filename
    coverage_path = candidate_directory / config.valid_day_coverage_filename
    candidate_points = read_parquet_file(candidate_path)
    valid_day_coverage = read_parquet_file(coverage_path)
    string_events, metric_audit = attach_candidate_metrics(
        confirmed_events,
        evidence,
        candidate_points,
        interval_minutes=config.interval_minutes,
        timezone=config.timezone,
    )
    history_config = config.history_config()
    event_profiles = attach_historical_features(
        string_events,
        valid_day_coverage,
        history_config,
    )
    string_profiles = build_string_profiles(
        event_profiles,
        valid_day_coverage,
        history_config,
    )
    unconfirmed = prepare_unconfirmed_runs(
        unconfirmed_source,
        timezone=config.timezone,
        interval_minutes=config.interval_minutes,
    )
    device_events = device_events.drop(
        columns=["target_member_numbers"], errors="ignore"
    )

    paths = {
        "confirmed_device_events": (
            output / "low_current_confirmed_device_events.parquet"
        ),
        "confirmed_memberships": (
            output / "low_current_confirmed_workbook_memberships.parquet"
        ),
        "event_crosswalk": output / "low_current_confirmed_event_crosswalk.parquet",
        "confirmed_string_events": (
            output / "low_current_confirmed_string_events.parquet"
        ),
        "confirmed_event_profiles": (
            output / "low_current_confirmed_event_profiles.parquet"
        ),
        "confirmed_string_profiles": (
            output / "low_current_confirmed_string_profiles.parquet"
        ),
        "unconfirmed_candidate_runs": (
            output / "low_current_unconfirmed_candidate_runs.parquet"
        ),
        "event_review_csv": output / "low_current_confirmed_event_review.csv",
        "string_profile_csv": output / "low_current_confirmed_string_profiles.csv",
        "device_event_csv": output / "low_current_confirmed_device_events.csv",
        "event_crosswalk_csv": output / "low_current_confirmed_event_crosswalk.csv",
        "unconfirmed_csv": output / "low_current_unconfirmed_candidate_runs.csv",
        "data_quality": output / "data_quality_summary.json",
        "config_snapshot": output / "config_snapshot.json",
        "summary": output / "summary.json",
        "workbook": output / "low_current_confirmed_profile_review.xlsx",
    }
    device_events.to_parquet(paths["confirmed_device_events"], index=False)
    memberships.to_parquet(paths["confirmed_memberships"], index=False)
    crosswalk.to_parquet(paths["event_crosswalk"], index=False)
    string_events.to_parquet(paths["confirmed_string_events"], index=False)
    event_profiles.to_parquet(paths["confirmed_event_profiles"], index=False)
    string_profiles.to_parquet(paths["confirmed_string_profiles"], index=False)
    unconfirmed.to_parquet(paths["unconfirmed_candidate_runs"], index=False)
    event_profiles.to_csv(
        paths["event_review_csv"], index=False, encoding="utf-8-sig"
    )
    string_profiles.to_csv(
        paths["string_profile_csv"], index=False, encoding="utf-8-sig"
    )
    device_events.to_csv(
        paths["device_event_csv"], index=False, encoding="utf-8-sig"
    )
    crosswalk.to_csv(
        paths["event_crosswalk_csv"], index=False, encoding="utf-8-sig"
    )
    unconfirmed.to_csv(
        paths["unconfirmed_csv"], index=False, encoding="utf-8-sig"
    )

    quality = {
        "replay_inputs": replay_audits,
        "workbook_crosswalk": crosswalk_audit,
        "candidate_metric_enrichment": metric_audit,
    }
    paths["data_quality"].write_text(
        json.dumps(quality, ensure_ascii=False, indent=2, default=_json_ready)
        + "\n",
        encoding="utf-8",
    )
    paths["config_snapshot"].write_text(
        json.dumps(config.to_dict(), ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    if not args.skip_workbook:
        frames = _workbook_frames(
            event_profiles=event_profiles,
            string_profiles=string_profiles,
            device_events=device_events,
            unconfirmed=unconfirmed,
            crosswalk=crosswalk,
            replay_audits=replay_audits,
            crosswalk_audit=crosswalk_audit,
            metric_audit=metric_audit,
            config=config,
        )
        _write_workbook(paths["workbook"], frames)

    manual_columns = ["manual_label", "manual_strings", "review_note"]
    manual_cells_filled = int(
        sum(
            device_events[column].fillna("").astype(str).str.strip().ne("").sum()
            for column in manual_columns
            if column in device_events
        )
    )
    summary = {
        "contract": {
            "profile_version": config.version,
            "source_algorithm_version": config.source_algorithm_version,
            "timezone": config.timezone,
            "interval_minutes": config.interval_minutes,
            "plant_ids": sorted(event_profiles["plant_id"].astype(str).unique()),
            "event_source": (
                "confirmed workbook ledger plus exact V1.7 Improved V2 events"
            ),
            "target_version_column": config.target_version_column,
            "candidates_are_event_filter": False,
            "unconfirmed_candidates_affect_classification": False,
            "causal_event_history_uses_confirmed_events_only": True,
            "manual_labels_applied": False,
            "provisional_thresholds": config.provisional_thresholds,
        },
        "counts": {
            "workbook_rows_all_versions": workbook_rows,
            "workbook_v2_device_events": int(len(device_events)),
            "workbook_v2_event_string_memberships": int(len(memberships)),
            "exact_confirmed_string_events": int(len(string_events)),
            "confirmed_physical_strings": int(
                string_events[["plant_id", "device_no", "string_no"]]
                .drop_duplicates()
                .shape[0]
            ),
            "unconfirmed_candidate_runs_auxiliary": int(len(unconfirmed)),
            "manual_workbook_cells_filled_not_applied": manual_cells_filled,
            "event_profile_classes": _classification_counts(event_profiles),
            "retrospective_string_classes": _classification_counts(
                string_profiles
            ),
            "fixed_time_string_profiles": int(
                string_profiles["fixed_time_tag"].sum()
            ),
        },
        "quality": quality,
        "configuration": config.to_dict(),
        "inputs": {
            "confirmed_events_workbook": str(args.confirmed_events_workbook),
            "replay_root": str(args.input_root),
            "candidate_points": str(candidate_path),
            "valid_day_coverage": str(coverage_path),
        },
        "outputs": {
            key: str(value)
            for key, value in paths.items()
            if key != "workbook" or not args.skip_workbook
        },
    }
    paths["summary"].write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, default=_json_ready)
        + "\n",
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2, default=_json_ready))


if __name__ == "__main__":
    main()
